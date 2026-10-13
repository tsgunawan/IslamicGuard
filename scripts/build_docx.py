"""T26 -- build manuscript/IslamicGuard_IJEEI_generated.docx from manuscript/draft.md on the IJEEI template.

The template supplies page size, margins, header/footer and the Normal style; its sample body
is removed. Layout follows the template's instructions (title 16 pt bold centred, section
headings bold capitals, table captions above, figure captions below, justified 10 pt body with
a 0.5 inch first-line indent). Citations stay in APA author-date form by the authors' decision
(DG-10); the template itself asks for numbered IEEE style.

Usage: python scripts/build_docx.py
"""
from __future__ import annotations

import re
from pathlib import Path

import docx
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent))
import omml  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
SRC, OUT, TEMPLATE = ROOT / "manuscript/draft.md", ROOT / "manuscript/IslamicGuard_IJEEI_generated.docx", ROOT / "Template IJEEI.docx"
# NOTE: manuscript/IslamicGuard_IJEEI.docx is now edited by hand (author photos, biographies, layout).
# This script therefore writes a separate "_generated" file and must never overwrite the edited one.
INLINE = re.compile(r"(\$[^$]+\$|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_runs(par, text, size=None, bold=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("$"):  # inline LaTeX -> native Word equation
            par._p.append(omml.omath(part[1:-1])); continue
        if part.startswith("**"):
            run = par.add_run(part[2:-2]); run.bold = True
        elif part.startswith("*"):
            run = par.add_run(part[1:-1]); run.italic = True
        elif part.startswith("`"):
            run = par.add_run(part[1:-1])
        else:
            run = par.add_run(part)
        if size:
            run.font.size = Pt(size)
        if bold:
            run.bold = True
    return par


def add_math(par, text, italic=True):
    """Italic run with _x / _{xyz} rendered as subscripts."""
    for piece in re.split(r"(_\{[^}]+\}|_[A-Za-zℓ])", text):
        if not piece:
            continue
        run = par.add_run(piece.lstrip("_").strip("{}") if piece.startswith("_") else piece)
        run.italic = italic
        if piece.startswith("_"):
            run.font.subscript = True


def _borders(table, spec):
    pr = table._tbl.tblPr
    el = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}"); b.set(qn("w:val"), spec.get(side, "nil")); b.set(qn("w:sz"), "8"); b.set(qn("w:space"), "0"); b.set(qn("w:color"), "000000")
        el.append(b)
    pr.append(el)


def no_borders(table):
    _borders(table, {})


def rule_borders(table):
    """Top and bottom rules only, the usual look of an algorithm float."""
    _borders(table, {"top": "single", "bottom": "single"})


def main():
    doc = docx.Document(TEMPLATE)
    body = doc.element.body
    for el in list(body):
        if not el.tag.endswith("sectPr"):
            body.remove(el)
    width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    lines = SRC.read_text(encoding="utf-8").splitlines()
    i, in_refs, after_title = 0, False, 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1; continue
        if ln.startswith("# "):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; add_runs(p, ln[2:], size=16, bold=True); after_title = 1
        elif ln.startswith("## "):
            in_refs = ln.strip() == "## REFERENCES"; after_title = 2
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); add_runs(p, ln[3:], bold=True)
        elif ln.startswith("### "):
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6); add_runs(p, ln[4:], bold=True)
        elif ln.startswith("!["):
            path = SRC.parent / re.search(r"\((.+?)\)", ln).group(1)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(path), width=min(width, Cm(15.4)))
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r, cells in enumerate(rows):
                for c, txt in enumerate(cells):
                    cp = t.cell(r, c).paragraphs[0]; add_runs(cp, txt, size=8, bold=(r == 0))
            if len(rows[0]) == 2:  # label/value table: narrow first column
                t.autofit = False
                for col, w in zip(t.columns, (int(width * 0.27), int(width * 0.73))):
                    col.width = w
                    for cell in col.cells:
                        cell.width = w
            if len(rows) <= 14:  # keep a short table on one page
                for r_ in t.rows[:-1]:
                    for cell in r_.cells:
                        for par in cell.paragraphs:
                            par.paragraph_format.keep_with_next = True
            doc.add_paragraph()
            continue
        elif ln.startswith("EQ: "):
            latex, num = re.match(r"EQ: (.*?)\s*\((\d+)\)$", ln).groups()
            t = doc.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
            no_borders(t)
            for col, cell, w in zip(t.columns, t.rows[0].cells, (int(width * 0.92), int(width * 0.08))):
                col.width = w; cell.width = w  # both are needed: grid column and cell width
            eq = t.cell(0, 0).paragraphs[0]; eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
            eq._p.append(omml.omath(latex))
            nump = t.cell(0, 1).paragraphs[0]; nump.alignment = WD_ALIGN_PARAGRAPH.RIGHT; nump.add_run(f"({num})")
            t.cell(0, 1).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        elif ln.strip() == "ALG-BEGIN":
            i += 1; rows = []
            while lines[i].strip() != "ALG-END":
                rows.append(lines[i]); i += 1
            t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
            rule_borders(t)
            cell = t.cell(0, 0); first = True
            for row in rows:
                p = cell.paragraphs[0] if first else cell.add_paragraph(); first = False
                indent = (len(row) - len(row.lstrip(" "))) // 2
                p.paragraph_format.left_indent = Cm(0.5 + 0.6 * indent); p.paragraph_format.first_line_indent = Cm(-0.5)
                p.paragraph_format.space_after = Pt(0)
                add_runs(p, row.strip(), size=9)
            doc.add_paragraph()
        elif re.match(r"\*\*(Table|Figure|Algorithm) \d+\.\*\*", ln):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; add_runs(p, ln.replace("**", ""), size=9)
            if not ln.startswith("**Figure"):
                p.paragraph_format.keep_with_next = True  # a table or algorithm caption stays with its body
        elif in_refs:
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.5); p.paragraph_format.first_line_indent = Inches(-0.5)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; add_runs(p, ln, size=9)
        else:
            p = doc.add_paragraph()
            if after_title == 1:  # author names, affiliations and emails, up to the abstract heading
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(p, ln, size=None if ln.startswith("**") else 9); i += 1; continue
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                if not ln.startswith(("**Keywords", "**Corresponding")):
                    p.paragraph_format.first_line_indent = Inches(0.5)
            add_runs(p, ln)
        i += 1
    doc.save(OUT)
    d = docx.Document(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}: {len(d.paragraphs)} paragraphs, {len(d.tables)} tables, {len(d.inline_shapes)} figures")


if __name__ == "__main__":
    main()
